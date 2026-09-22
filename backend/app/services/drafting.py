"""Organization-scoped legal drafting workflows."""

from __future__ import annotations

from io import BytesIO
from uuid import UUID

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from ..config import Settings
from ..core.drafting import generate_clause
from ..core.generation import AnswerGenerator
from ..core.risk import score_document_risk
from ..models import (
    Draft,
    DraftClause,
    DraftClauseSource,
    DraftStatus,
    DraftTemplate,
    DraftVersion,
    Organization,
    User,
)
from ..schemas import (
    DraftClauseResponse,
    DraftDetail,
    DraftSummary,
    DraftTemplateResponse,
)
from .documents import get_or_create_organization

_BUILTIN_TEMPLATES = (
    {
        "slug": "nda",
        "name": "Non-Disclosure Agreement",
        "document_type": "NDA",
        "input_schema": {
            "disclosing_party": {"label": "Disclosing party", "required": True},
            "receiving_party": {"label": "Receiving party", "required": True},
            "purpose": {"label": "Purpose", "required": True},
            "term": {"label": "Term", "required": True},
        },
        "clause_outline": [
            {"heading": "Definitions", "purpose": "Define confidential information."},
            {"heading": "Confidentiality Obligations", "purpose": "Limit use and disclosure."},
            {"heading": "Permitted Disclosures", "purpose": "Describe required and permitted disclosures."},
            {"heading": "Term and Survival", "purpose": "Set the agreement and confidentiality periods."},
            {"heading": "Governing Law", "purpose": "Set a governing law placeholder."},
        ],
    },
    {
        "slug": "msa",
        "name": "Master Services Agreement",
        "document_type": "MSA",
        "input_schema": {
            "client": {"label": "Client", "required": True},
            "provider": {"label": "Service provider", "required": True},
            "services": {"label": "Services", "required": True},
            "fees": {"label": "Fees", "required": True},
        },
        "clause_outline": [
            {"heading": "Services", "purpose": "Describe the services and delivery standard."},
            {"heading": "Fees and Payment", "purpose": "Set invoicing and payment terms."},
            {"heading": "Confidentiality", "purpose": "Protect business information."},
            {"heading": "Limitation of Liability", "purpose": "Set a commercially balanced liability cap."},
            {"heading": "Termination", "purpose": "Set termination rights and effects."},
        ],
    },
    {
        "slug": "sow",
        "name": "Statement of Work",
        "document_type": "SOW",
        "input_schema": {
            "customer": {"label": "Customer", "required": True},
            "supplier": {"label": "Supplier", "required": True},
            "deliverables": {"label": "Deliverables", "required": True},
            "deadline": {"label": "Deadline", "required": True},
        },
        "clause_outline": [
            {"heading": "Scope", "purpose": "Define the agreed work."},
            {"heading": "Deliverables and Acceptance", "purpose": "Define delivery and acceptance."},
            {"heading": "Schedule", "purpose": "Set milestones and deadline."},
            {"heading": "Fees", "purpose": "Set the project price and payment."},
        ],
    },
    {
        "slug": "employment-agreement",
        "name": "Employment Agreement",
        "document_type": "Employment Agreement",
        "input_schema": {
            "employer": {"label": "Employer", "required": True},
            "employee": {"label": "Employee", "required": True},
            "role": {"label": "Role", "required": True},
            "salary": {"label": "Salary", "required": True},
        },
        "clause_outline": [
            {"heading": "Appointment and Duties", "purpose": "Define the role and duties."},
            {"heading": "Compensation", "purpose": "Set salary and payment terms."},
            {"heading": "Confidentiality and IP", "purpose": "Protect confidential information and work product."},
            {"heading": "Termination", "purpose": "Set termination rights."},
        ],
    },
    {
        "slug": "ip-assignment",
        "name": "IP Assignment Agreement",
        "document_type": "IP Assignment",
        "input_schema": {
            "assignor": {"label": "Assignor", "required": True},
            "assignee": {"label": "Assignee", "required": True},
            "work": {"label": "Work or invention", "required": True},
            "consideration": {"label": "Consideration", "required": True},
        },
        "clause_outline": [
            {"heading": "Assignment", "purpose": "Assign the identified intellectual property."},
            {"heading": "Further Assurances", "purpose": "Require cooperation for perfection."},
            {"heading": "Representations", "purpose": "Set ownership and authority representations."},
            {"heading": "Consideration", "purpose": "Record the consideration."},
        ],
    },
)


async def list_templates(session: AsyncSession, settings: Settings) -> list[DraftTemplateResponse]:
    organization = await get_or_create_organization(session, settings.default_org_id)
    await _ensure_builtins(session, organization.id)
    result = await session.execute(
        select(DraftTemplate).where(DraftTemplate.organization_id == organization.id).order_by(DraftTemplate.name)
    )
    return [_template_response(template) for template in result.scalars().all()]


async def create_draft(
    session: AsyncSession,
    settings: Settings,
    *,
    template_id: UUID,
    title: str,
    inputs: dict[str, str],
    user: User,
    generator: AnswerGenerator,
) -> DraftDetail | None:
    organization = await get_or_create_organization(session, settings.default_org_id)
    await _ensure_builtins(session, organization.id)
    template = await session.scalar(
        select(DraftTemplate).where(
            DraftTemplate.id == template_id, DraftTemplate.organization_id == organization.id
        )
    )
    if template is None:
        return None
    _validate_inputs(template, inputs)
    draft = Draft(
        organization_id=organization.id,
        user_id=user.id,
        template_id=template.id,
        title=title,
        inputs=inputs,
        risk_breakdown={},
    )
    session.add(draft)
    await session.flush()
    await _generate_all(session, draft, template, inputs, generator)
    _snapshot(draft, user)
    await session.flush()
    return _detail(draft)


async def list_drafts(session: AsyncSession, settings: Settings, user: User) -> list[DraftSummary]:
    organization = await get_or_create_organization(session, settings.default_org_id)
    result = await session.execute(
        select(Draft).where(
            Draft.organization_id == organization.id, Draft.user_id == user.id
        ).order_by(Draft.updated_at.desc())
    )
    return [_summary(draft) for draft in result.scalars().all()]


async def get_draft(session: AsyncSession, settings: Settings, draft_id: UUID, user: User) -> DraftDetail | None:
    draft = await _load(session, settings, draft_id, user)
    return _detail(draft) if draft else None


async def update_clause(
    session: AsyncSession, settings: Settings, draft_id: UUID, clause_id: UUID,
    *, heading: str, body: str, user: User,
) -> DraftDetail | None:
    draft = await _load(session, settings, draft_id, user)
    if draft is None:
        return None
    clause = next((item for item in draft.clauses if item.id == clause_id), None)
    if clause is None:
        return None
    clause.heading = heading
    clause.body = body
    clause.source = DraftClauseSource.EDITED
    clause.version += 1
    _refresh_risk(draft)
    _snapshot(draft, user)
    await session.flush()
    return _detail(draft)


async def regenerate_clause(
    session: AsyncSession, settings: Settings, draft_id: UUID, clause_id: UUID,
    *, instruction: str | None, user: User, generator: AnswerGenerator,
) -> DraftDetail | None:
    draft = await _load(session, settings, draft_id, user)
    if draft is None:
        return None
    clause = next((item for item in draft.clauses if item.id == clause_id), None)
    if clause is None:
        return None
    template = draft.template
    generated = await generate_clause(
        generator,
        template_name=template.name,
        clause_heading=clause.heading,
        clause_outline=template.clause_outline,
        inputs=draft.inputs,
        prior_clauses=[item.body for item in draft.clauses if item.ordinal < clause.ordinal],
        instruction=instruction,
    )
    clause.heading = generated.heading
    clause.body = generated.body
    clause.rationale = generated.rationale
    clause.risk_notes = generated.risk_notes
    clause.source = DraftClauseSource.GENERATED
    clause.version += 1
    _refresh_risk(draft)
    _snapshot(draft, user)
    await session.flush()
    return _detail(draft)


async def delete_draft(session: AsyncSession, settings: Settings, draft_id: UUID, user: User) -> bool:
    draft = await _load(session, settings, draft_id, user)
    if draft is None:
        return False
    await session.delete(draft)
    await session.flush()
    return True


def export_docx(draft: DraftDetail) -> bytes:
    document = DocxDocument()
    document.add_heading(draft.title, level=0)
    for clause in draft.clauses:
        document.add_heading(clause.heading, level=1)
        document.add_paragraph(clause.body)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


async def _ensure_builtins(session: AsyncSession, organization_id: UUID) -> None:
    existing = set((await session.scalars(select(DraftTemplate.slug).where(DraftTemplate.organization_id == organization_id))).all())
    for definition in _BUILTIN_TEMPLATES:
        if definition["slug"] not in existing:
            session.add(DraftTemplate(organization_id=organization_id, is_builtin=True, **definition))
    await session.flush()


async def _generate_all(session: AsyncSession, draft: Draft, template: DraftTemplate, inputs: dict[str, str], generator: AnswerGenerator) -> None:
    prior: list[str] = []
    for ordinal, outline in enumerate(template.clause_outline):
        generated = await generate_clause(
            generator,
            template_name=template.name,
            clause_heading=str(outline["heading"]),
            clause_outline=template.clause_outline,
            inputs=inputs,
            prior_clauses=prior,
        )
        draft.clauses.append(DraftClause(
            organization_id=draft.organization_id,
            ordinal=ordinal,
            heading=generated.heading,
            body=generated.body,
            rationale=generated.rationale,
            risk_notes=generated.risk_notes,
            source=DraftClauseSource.GENERATED,
        ))
        prior.append(generated.body)
    _refresh_risk(draft)


def _refresh_risk(draft: Draft) -> None:
    assessment = score_document_risk("\n\n".join(clause.body for clause in draft.clauses))
    draft.risk_breakdown = {"score": assessment.score, **assessment.breakdown}


def _snapshot(draft: Draft, user: User) -> None:
    version = len(draft.versions) + 1
    draft.versions.append(DraftVersion(
        organization_id=draft.organization_id,
        version=version,
        created_by=user.id,
        snapshot={"title": draft.title, "clauses": [
            {"heading": clause.heading, "body": clause.body, "source": clause.source.value}
            for clause in draft.clauses
        ]},
    ))


async def _load(session: AsyncSession, settings: Settings, draft_id: UUID, user: User) -> Draft | None:
    organization = await get_or_create_organization(session, settings.default_org_id)
    result = await session.execute(
        select(Draft).options(selectinload(Draft.template), selectinload(Draft.clauses), selectinload(Draft.versions)).where(
            Draft.id == draft_id, Draft.organization_id == organization.id, Draft.user_id == user.id
        )
    )
    return result.scalar_one_or_none()


def _validate_inputs(template: DraftTemplate, inputs: dict[str, str]) -> None:
    required = [key for key, definition in template.input_schema.items() if definition.get("required")]
    missing = [key for key in required if not inputs.get(key, "").strip()]
    if missing:
        raise ValueError(f"Missing required inputs: {', '.join(missing)}")


def _template_response(template: DraftTemplate) -> DraftTemplateResponse:
    return DraftTemplateResponse.model_validate(template, from_attributes=True)


def _summary(draft: Draft) -> DraftSummary:
    return DraftSummary.model_validate(draft, from_attributes=True)


def _detail(draft: Draft) -> DraftDetail:
    return DraftDetail(
        id=draft.id, title=draft.title, template_id=draft.template_id, status=draft.status,
        updated_at=draft.updated_at, inputs=draft.inputs, risk_breakdown=draft.risk_breakdown,
        clauses=[DraftClauseResponse.model_validate(clause, from_attributes=True) for clause in draft.clauses],
        version=len(draft.versions),
    )
