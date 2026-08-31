"""Extract searchable text from supported legal document formats."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

SUPPORTED_EXTENSIONS = frozenset({".csv", ".docx", ".pdf", ".txt", ".xlsx"})


def load_document(file_path: str | Path) -> str:
    """Load a document and return non-empty text suitable for processing."""
    path = Path(file_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    loaders = {
        ".csv": _load_csv,
        ".docx": _load_docx,
        ".pdf": _load_pdf,
        ".txt": _load_txt,
        ".xlsx": _load_xlsx,
    }
    try:
        loader = loaders[path.suffix.lower()]
    except KeyError as exc:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(
            f"Unsupported file type '{path.suffix.lower()}'. Supported types: {supported}"
        ) from exc

    text = loader(path).strip()
    if not text:
        raise ValueError(f"No text could be extracted from '{path.name}'")
    return text


def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _load_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.reader(csv_file)
        rows = list(reader)
    if not rows:
        return ""
    return _format_table(path.stem, rows[0], rows[1:])


def _load_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue
            sections.append(_format_table(worksheet.title, rows[0], rows[1:]))
    finally:
        workbook.close()
    return "\n\n".join(sections)


def _load_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    return "\n\n".join(paragraph for paragraph in paragraphs if paragraph)


def _load_pdf(path: Path) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                pages.append(f"Page {page_number}\n{page_text}")
    return "\n\n".join(pages)


def _format_table(
    section_name: str,
    raw_headers: Iterable[Any],
    rows: Iterable[Iterable[Any]],
) -> str:
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    lines = [f"Table: {section_name}"]
    for row_number, row in enumerate(rows, start=2):
        values = list(row)
        fields = [
            f"{header or f'Column {column_number}'}: {value}"
            for column_number, (header, value) in enumerate(
                zip(headers, values, strict=False), start=1
            )
            if value is not None and str(value).strip()
        ]
        if fields:
            lines.append(f"Row {row_number} | " + " | ".join(fields))
    return "\n".join(lines)